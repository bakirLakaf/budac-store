"""
اختبارات انحدار — التدفّقات العادية لازم تبقى تخدم بعد الإصلاحات.

⚠️ هاد الملف أهم من اختبارات الأمان بمعنى واحد: الإصلاح الأمني
اللي يكسر البيع أسوأ من الثغرة. كل تدفّق يجيب فلوس ولا يستعملو
الأدمن يوميًا لازم يكون هنا.
"""

import io

import pytest

from models import db, Product, Order, Review, SiteVisit


# ══════════════════════════════════════════════════════════════
# صفحات المتجر
# ══════════════════════════════════════════════════════════════

@pytest.mark.parametrize("path", [
    "/", "/cart", "/track", "/my-orders", "/search?q=شريط",
    "/sitemap.xml", "/robots.txt", "/admin/login",
])
def test_public_pages_render(client, seed, path):
    assert client.get(path).status_code == 200


def test_product_page_renders(client, seed):
    assert client.get(f"/product/{seed['product_slug']}").status_code == 200


def test_landing_page_renders(client, seed):
    assert client.get(f"/lp/{seed['page_slug']}").status_code == 200


def test_homepage_category_filter_still_works(client, app, seed):
    """
    ⚠️ الفلترة بالتصنيف ولّات من السيرفر (?category=) بدل الفلترة
    الفورية بـJavaScript — هاذي ماتت مع الترقيم لأن المنتجات ماشي
    كلها فالصفحة. لازم تبقى تخدم للروابط المباشرة وللزواحف.
    """
    assert client.get("/").status_code == 200
    assert client.get("/?category=decor").status_code == 200

    body = client.get("/?category=decor").get_data(as_text=True)
    assert "شريط كريستال" in body


def test_homepage_stays_bounded_on_a_huge_catalog(app, seed):
    """
    ⚠️ هاد الاختبار تبدّل مرتين مع الطراز، وهذا مقصود — يقيس
    السلوك اللي يهمّ حقًّا، ماشي تفاصيل التنفيذ:

      v139: استعلامان يحمّلو الكتالوج **كامل** مرتين
      v141: ترقيم من السيرفر — 9 منتجات، لكن كل نقرة إعادة تحميل
      v142: ترقيم فالمتصفح — الكتالوج كامل مرة وحدة، تبديل فوري
            **مع حارس**: فوق MAX_CLIENT_PAGINATION نرجعو للسيرفر

    اللي نحميوه درك: كتالوج ضخم ما يخصّوش يبعث كلشي للمتصفح.
    """
    from models import Product
    from routes.shop import MAX_CLIENT_PAGINATION, PRODUCTS_PER_PAGE
    from flask import template_rendered

    with app.app_context():
        for i in range(MAX_CLIENT_PAGINATION + 20):
            db.session.add(Product(
                name=f"منتج ضخم {i}", slug=f"bulk-{i}",
                price=100, is_active=True,
            ))
        db.session.commit()
        total = Product.query.filter_by(is_active=True).count()

    captured = []

    def record(sender, template, context, **extra):
        if template.name == "shop/home.html":
            captured.append(context)

    template_rendered.connect(record, app)
    try:
        app.test_client().get("/")
    finally:
        template_rendered.disconnect(record, app)

    context = captured[0]
    assert total > MAX_CLIENT_PAGINATION
    assert context["client_paginated"] is False, "الحارس ما اشتغلش"
    assert len(context["products"]) == PRODUCTS_PER_PAGE, (
        "كتالوج ضخم مازال يتبعث كامل للمتصفح"
    )


def test_small_catalog_is_sent_whole_for_instant_switching(app, seed):
    """
    الوجه الثاني: كتالوج صغير (الوضع العادي) يتبعث كامل — هو
    اللي يخلّي تبديل الصفحات فوريًا بلا إعادة تحميل.
    """
    from models import Product
    from flask import template_rendered

    with app.app_context():
        for i in range(30):
            db.session.add(Product(
                name=f"منتج {i}", slug=f"small-{i}", price=100, is_active=True,
            ))
        db.session.commit()
        total = Product.query.filter_by(is_active=True).count()

    captured = []

    def record(sender, template, context, **extra):
        if template.name == "shop/home.html":
            captured.append(context)

    template_rendered.connect(record, app)
    try:
        app.test_client().get("/")
    finally:
        template_rendered.disconnect(record, app)

    context = captured[0]
    assert context["client_paginated"] is True
    assert len(context["products"]) == total


# ══════════════════════════════════════════════════════════════
# السلة والشراء
# ══════════════════════════════════════════════════════════════

def test_full_checkout_flow(client, app, seed):
    """التدفّق الكامل: زيد للسلة → اشتري → صفحة النجاح."""
    add = client.post("/cart/add", data={
        "kind": "product", "item_id": seed["product_id"], "quantity": "2",
    })
    assert add.status_code in (200, 302)

    cart = client.get("/api/cart").get_json()
    # ⚠️ "count" = عدد الأسطر فالسلة، ماشي عدد القطع
    assert cart["count"] == 1
    assert cart["lines"][0]["quantity"] == 2
    assert cart["subtotal"] == 2400

    checkout = client.post("/checkout", data={
        "name": "علي بن علي", "phone": "0666255200",
        "wilaya_code": "39", "commune": "الوادي",
        "carrier_code": "zrexpress", "delivery_type": "home",
    })
    assert checkout.status_code == 302

    with app.app_context():
        order = Order.query.first()
        assert order is not None
        assert order.total == 3000          # 2400 + 600
        number = order.order_number

    # صاحب الطلبية يشوف صفحة النجاح
    assert client.get(f"/order/success/{number}").status_code == 200


def test_cart_update_and_remove(client, seed):
    client.post("/cart/add", data={
        "kind": "product", "item_id": seed["product_id"], "quantity": "1",
    })
    cart = client.get("/api/cart").get_json()
    key = cart["lines"][0]["key"]

    client.post("/api/cart/update", json={"key": key, "quantity": 3})
    assert client.get("/api/cart").get_json()["lines"][0]["quantity"] == 3

    client.post("/api/cart/update", json={"key": key, "quantity": 0})
    assert client.get("/api/cart").get_json()["count"] == 0


def test_shipping_price_api(client, seed):
    response = client.get(
        "/api/shipping-price?carrier=zrexpress&wilaya=39&type=home"
    )
    assert response.status_code == 200
    data = response.get_json()
    assert data["price"] == 600, data


# ══════════════════════════════════════════════════════════════
# تتبّع الطلبية وإلغاؤها
# ══════════════════════════════════════════════════════════════

def _place_order(client, seed):
    client.post("/cart/add", data={
        "kind": "product", "item_id": seed["product_id"], "quantity": "1",
    })
    client.post("/checkout", data={
        "name": "علي بن علي", "phone": "0666255200",
        "wilaya_code": "39", "commune": "الوادي",
        "carrier_code": "zrexpress", "delivery_type": "home",
    })


def test_order_tracking_with_correct_credentials(client, app, seed):
    _place_order(client, seed)
    with app.app_context():
        number = Order.query.first().order_number

    body = client.post("/track", data={
        "order_number": number, "phone": "0666255200",
    }, follow_redirects=True).get_data(as_text=True)
    assert number in body


def test_order_tracking_rejects_wrong_phone(client, app, seed):
    _place_order(client, seed)
    with app.app_context():
        number = Order.query.first().order_number

    body = client.post("/track", data={
        "order_number": number, "phone": "0555111222",
    }, follow_redirects=True).get_data(as_text=True)
    assert number not in body


def test_customer_can_cancel_pending_order(client, app, seed):
    _place_order(client, seed)
    with app.app_context():
        number = Order.query.first().order_number

    client.post(f"/my-orders/{number}/cancel",
                data={"phone": "0666255200"}, follow_redirects=True)

    with app.app_context():
        assert Order.query.first().status == Order.STATUS_CANCELLED


def test_customer_cannot_cancel_someone_elses_order(client, app, seed):
    _place_order(client, seed)
    with app.app_context():
        number = Order.query.first().order_number

    client.post(f"/my-orders/{number}/cancel",
                data={"phone": "0555999888"}, follow_redirects=True)

    with app.app_context():
        assert Order.query.first().status == Order.STATUS_PENDING


# ══════════════════════════════════════════════════════════════
# التقييمات
# ══════════════════════════════════════════════════════════════

def test_review_submission_and_moderation(client, app, seed):
    """التقييم يتسجّل غير معتمد — ما يبانش قبل موافقة الأدمن."""
    client.post("/reviews", data={
        "author_name": "فاطمة", "rating": "5", "comment": "خدمة مليحة بزاف",
    }, follow_redirects=True)

    with app.app_context():
        review = Review.query.first()
        assert review is not None
        assert review.is_approved is False


# ══════════════════════════════════════════════════════════════
# لوحة التحكم
# ══════════════════════════════════════════════════════════════

def test_admin_login_and_dashboard(client, seed):
    response = client.post("/admin/login", data={
        "username": seed["admin_username"],
        "password": seed["admin_password"],
    }, follow_redirects=True)
    assert response.status_code == 200
    assert client.get("/admin/").status_code == 200


def test_admin_login_rejects_wrong_password(client, seed):
    client.post("/admin/login", data={
        "username": seed["admin_username"], "password": "wrong",
    })
    assert client.get("/admin/").status_code == 302


@pytest.mark.parametrize("path", [
    "/admin/", "/admin/orders", "/admin/products", "/admin/coupons",
    "/admin/reviews", "/admin/shipping", "/admin/settings",
])
def test_admin_pages_render_for_logged_in_admin(logged_in_admin, path):
    assert logged_in_admin.get(path).status_code == 200


def test_admin_can_change_order_status(logged_in_admin, app, seed):
    with app.app_context():
        order = Order(
            order_number="BD-40004000", customer_name="علي",
            customer_phone="0666255200", wilaya_code="39", wilaya_name="الوادي",
            subtotal=1200, shipping_price=600, total=1800,
            status=Order.STATUS_PENDING,
        )
        db.session.add(order)
        db.session.commit()
        order_id = order.id

    logged_in_admin.post(f"/admin/orders/{order_id}/status",
                         data={"status": Order.STATUS_CONFIRMED},
                         follow_redirects=True)

    with app.app_context():
        assert db.session.get(Order, order_id).status == Order.STATUS_CONFIRMED


def test_admin_action_is_written_to_audit_log(logged_in_admin, app, seed):
    """سجل التدقيق يسجّل الدخول فعلاً."""
    from models import AdminAction

    with app.app_context():
        logins = AdminAction.query.filter_by(action=AdminAction.LOGIN).count()
        assert logins >= 1


# ══════════════════════════════════════════════════════════════
# Excel
# ══════════════════════════════════════════════════════════════

def test_excel_export_then_import_roundtrip(logged_in_admin, app, seed):
    """
    ⚠️ بدّلنا الاستيراد لوضع read_only — القراءة بالبثّ ما تدعمش
    sheet[1]. هاد الاختبار يتأكد أن الدورة الكاملة مازالت تخدم.
    """
    export = logged_in_admin.get("/admin/products/export")
    assert export.status_code == 200
    assert len(export.data) > 1000

    upload = logged_in_admin.post(
        "/admin/products/import",
        data={"file": (io.BytesIO(export.data), "products.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert upload.status_code == 200
    assert "فشل الاستيراد" not in upload.get_data(as_text=True)


def test_excel_export_neutralises_formula_injection(app, seed):
    """
    اسم منتج يبدا بـ= يتنفّذ كصيغة كي يفتح الأدمن الملف.
    الأسماء تجي من استيراد YouCan — بيانات خارجية.
    """
    from openpyxl import load_workbook
    from services.excel_io import export_workbook

    with app.app_context():
        product = db.session.get(Product, seed["product_id"])
        product.name = '=HYPERLINK("http://evil.com","اضغط")'
        db.session.commit()

        stream = export_workbook()

    workbook = load_workbook(stream)
    found = False
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "HYPERLINK" in cell.value:
                    found = True
                    assert cell.value.startswith("'"), "الصيغة ما اتحيّدتش"
    assert found, "ما لقيناش الخانة — الاختبار ما فحص والو"


def test_safe_cell_keeps_numbers_intact():
    """
    ⚠️ الحماية تخصّ النصوص برك — لو حوّلنا الأرقام لنصوص،
    تنكسر الحسابات والفرز فالملف وينكسر الاستيراد.
    """
    from services.excel_io import _safe_cell

    assert _safe_cell(1200) == 1200
    assert _safe_cell(None) is None
    assert _safe_cell(True) is True
    assert _safe_cell("شريط عادي") == "شريط عادي"
    assert _safe_cell("=1+1") == "'=1+1"
    assert _safe_cell("@SUM(A1)") == "'@SUM(A1)"


# ══════════════════════════════════════════════════════════════
# تتبّع الزوار
# ══════════════════════════════════════════════════════════════

def test_spoofed_forwarded_header_cannot_flood_visits(client, app, seed):
    """
    🔴 كانت analytics تقرا X-Forwarded-For خامًا وتاخذ **أول** قيمة —
    واللي هي بالضبط الجزء اللي يكتبو العميل. فمهاجم يبدّلها فكل طلب
    → بصمة جديدة → صف جديد → تضخّم جدول بلا حد.

    ⚠️ نحاكيو البروكسي كيما يخدم فعلاً: Railway **يزيد** العنوان
    الحقيقي اللي شافو فآخر السلسلة. ProxyFix(x_for=1) ياخذ آخر
    قيمة — يعني الحقيقية — ويحطّها فـremote_addr.

    فالمهاجم يقدر يكتب اللي يحب فبداية السلسلة، وما يتبدّل والو.
    """
    real_client_ip = "197.200.10.55"

    for i in range(30):
        client.get("/", headers={
            # الجزء الأول من العميل (مزوّر) + العنوان الحقيقي اللي
            # يزيدو البروكسي فالآخر
            "X-Forwarded-For": f"10.0.0.{i}, {real_client_ip}",
        })

    with app.app_context():
        assert SiteVisit.query.count() <= 2, (
            "الهيدر المزوّر مازال يولّد صفوفًا — البصمة تتبدّل مع كل طلب"
        )


def test_visitor_identity_uses_proxyfix_not_raw_header(app, seed):
    """
    ⚠️ اختبار للسبب الجذري: نتأكدو أن الكود ما يقراش الهيدر الخام.

    مهم يتوثّق: أمان هاد الجزء **مشروط** بأن `x_for` فـProxyFix
    يطابق عدد البروكسيات الحقيقية. Railway = بروكسي واحد = x_for=1.
    لو زدت CDN قدّامو (Cloudflare مثلاً)، لازم x_for=2 — وإلا
    ترجع الثغرة من باب آخر.
    """
    import inspect
    from services import analytics

    source = inspect.getsource(analytics._visitor_key)
    assert "X-Forwarded-For" not in source or "request.remote_addr" in source
    assert "request.remote_addr" in source

    # وProxyFix مركّب فعلاً بعدد بروكسي واحد (Railway)
    import app as app_module
    app_source = inspect.getsource(app_module.create_app)
    assert "ProxyFix" in app_source
    assert "x_for=1" in app_source


def test_visits_are_still_counted(client, app, seed):
    """⚠️ الإصلاح ما يخصّوش يعطّل العدّاد كليًا."""
    client.get("/")
    with app.app_context():
        assert SiteVisit.query.count() >= 1


# ══════════════════════════════════════════════════════════════
# تحديد معدّل الطلبات
# ══════════════════════════════════════════════════════════════

def test_rate_limits_are_actually_registered(app):
    """
    ⚠️ الحدود تتطبّق بتبديل app.view_functions بعد تسجيل المسارات.
    لو الآلية ما خدماتش، الحدود تبان مكتوبة وما تشتغل والو —
    وهذا أسوأ من غيابها (أمان وهمي).
    """
    limiter = app.extensions.get("budac_limiter")
    assert limiter is not None, "المحدِّد ما اتركّبش"

    limited = [
        endpoint for endpoint in app.view_functions
        if endpoint in {
            "cart.checkout", "landing.submit", "landing.thanks",
            "shop.add_store_review", "cart.my_orders_lookup", "shop.search",
        }
    ]
    assert len(limited) >= 6, f"مسارات ناقصة: {limited}"

    assert limiter.enabled


def test_rate_limit_blocks_excessive_review_spam(app, seed):
    """5 تقييمات فالساعة — السادس يترفض بـ429."""
    app.config["WTF_CSRF_ENABLED"] = False
    client = app.test_client()

    codes = []
    for i in range(9):
        response = client.post("/reviews", data={
            "author_name": f"زائر {i}", "rating": "5",
            "comment": "تعليق تجريبي طويل شوية",
        })
        codes.append(response.status_code)

    assert 429 in codes, f"الحد ما اشتغلش: {codes}"


# ══════════════════════════════════════════════════════════════
# السابق/التالي للطلبيات
# ══════════════════════════════════════════════════════════════

def test_order_neighbours_navigation(logged_in_admin, app, seed):
    """صفحة الطلبية تعرض روابط السابق/التالي صحيحة حسب الترتيب."""
    from datetime import datetime, timedelta

    with app.app_context():
        base = datetime.utcnow()
        order_ids = []
        for i in range(3):
            o = Order(
                order_number=f"BD-NAV{i:03d}", customer_name=f"زبون {i}",
                customer_phone="0555000000", wilaya_code="16",
                wilaya_name="الجزائر", subtotal=500, shipping_price=0,
                total=500, status=Order.STATUS_PENDING,
                created_at=base - timedelta(minutes=i),
            )
            db.session.add(o)
            db.session.flush()
            order_ids.append(o.id)
        db.session.commit()

    # الطلبية الوسطى (الثانية بالترتيب الزمني تنازليًا)
    middle_id = order_ids[1]
    resp = logged_in_admin.get(f"/admin/orders/{middle_id}")
    assert resp.status_code == 200
    html = resp.get_data(as_text=True)

    assert "neighbours-pos" in html
    assert f"BD-NAV{2:03d}" in html or f"BD-NAV{0:03d}" in html


# ══════════════════════════════════════════════════════════════
# تحديث حالة التتبّع من ZR — ما يطيحش بصفحة 500
# ══════════════════════════════════════════════════════════════

def test_refresh_tracking_handles_unexpected_zr_response(logged_in_admin, app, seed):
    """
    استثناء غير متوقع من get_parcel_status (مو ZRError) لازم يرجّع
    رسالة واضحة، ماشي صفحة 500 خام.

    ⚠️ هاد الاختبار وُجد بعد بلاغ حقيقي: زر "حدّث حالة التتبّع من ZR"
    كان يوقع صفحة 500 عامة بلا أي تفاصيل تفيد الأدمن.
    """
    from unittest.mock import patch

    with app.app_context():
        order = Order(
            order_number="BD-TRK001", customer_name="زبون تتبّع",
            customer_phone="0555000000", wilaya_code="16",
            wilaya_name="الجزائر", subtotal=500, shipping_price=0,
            total=500, status=Order.STATUS_CONFIRMED,
            tracking_number="FCQIKGY725-ZR-24",
        )
        db.session.add(order)
        db.session.commit()
        oid = order.id

    # نحاكيو استثناء غير متوقع (ماشي ZRError) — بالضبط نوع العطب
    # اللي كان يطيح 500
    with patch(
        "services.zrexpress.get_parcel_status",
        side_effect=AttributeError("'NoneType' object has no attribute 'get'"),
    ):
        resp = logged_in_admin.post(
            f"/admin/orders/{oid}/refresh-tracking", follow_redirects=True
        )

    assert resp.status_code == 200, "لازم يرجّع رسالة واضحة ماشي 500"
    html = resp.get_data(as_text=True)
    assert "خطأ غير متوقع" in html


def test_refresh_tracking_handles_malformed_status_response(logged_in_admin, app, seed):
    """رد ZR بشكل غير متوقع (None ولا بلا مفتاح state) يرجّع رسالة واضحة."""
    from unittest.mock import patch

    with app.app_context():
        order = Order(
            order_number="BD-TRK002", customer_name="زبون تتبّع 2",
            customer_phone="0555000001", wilaya_code="16",
            wilaya_name="الجزائر", subtotal=500, shipping_price=0,
            total=500, status=Order.STATUS_CONFIRMED,
            tracking_number="FCQIKGY726-ZR-25",
        )
        db.session.add(order)
        db.session.commit()
        oid = order.id

    with patch("services.zrexpress.get_parcel_status", return_value=None):
        resp = logged_in_admin.post(
            f"/admin/orders/{oid}/refresh-tracking", follow_redirects=True
        )

    assert resp.status_code == 200
    html = resp.get_data(as_text=True)
    assert "ما فيهش حالة واضحة" in html


# ══════════════════════════════════════════════════════════════
# get_parcel_status: ZR بدّلت شكل حقل state لكائن — ماشي نص
# ══════════════════════════════════════════════════════════════

def test_get_parcel_status_handles_new_zr_state_object(monkeypatch):
    """
    بلاغ إنتاج حقيقي (v148): ZR غيّرت رد /parcels/{id} — حقل `state`
    ولّى كائن {id, name, description, color, ...} بدل نص بسيط.

    الرد الحقيقي اللي جبناه من logs الإنتاج:
        {"state": {"id": "...", "name": "encaisse",
                    "description": "Encaissé", "color": "00796b", ...}}

    كان يخلّي `order.carrier_state = status["state"]` يحاول يحفظ
    dict كامل فعمود نص قصير، وpsycopg2 يرفض بـ:
        ProgrammingError: can't adapt type 'dict'
    """
    from unittest.mock import patch, MagicMock
    from services.zrexpress import get_parcel_status

    monkeypatch.setenv("ZR_TENANT_ID", "test-tenant")
    monkeypatch.setenv("ZR_API_KEY", "test-key")

    fake_response = MagicMock()
    fake_response.status_code = 200
    fake_response.content = b"{}"
    fake_response.headers = {}
    fake_response.json.return_value = {
        "state": {
            "id": "f34dfc2c-ecfe-4e06-9a19-430d2d0b27b9",
            "name": "encaisse",
            "description": "Encaissé",
            "isBlocking": False,
            "isLocked": False,
            "visibleFor": 15,
            "editableBy": 2,
            "color": "00796b",
        }
    }

    with patch("services.zrexpress.requests.request", return_value=fake_response):
        result = get_parcel_status("FCQIKGY725-ZR-24")

    assert isinstance(result["state"], str), "state لازم يكون نص، ماشي dict"
    assert result["state"] == "Encaissé"
    assert result["state_id"] == "f34dfc2c-ecfe-4e06-9a19-430d2d0b27b9"


def test_get_parcel_status_still_supports_old_plain_string_format(monkeypatch):
    """الشكل القديم (state = نص مباشر) لازم يبقى يخدم، ماشي غير الجديد."""
    from unittest.mock import patch, MagicMock
    from services.zrexpress import get_parcel_status

    monkeypatch.setenv("ZR_TENANT_ID", "test-tenant")
    monkeypatch.setenv("ZR_API_KEY", "test-key")

    fake_response = MagicMock()
    fake_response.status_code = 200
    fake_response.content = b"{}"
    fake_response.headers = {}
    fake_response.json.return_value = {"state": "delivered", "stateId": "xyz"}

    with patch("services.zrexpress.requests.request", return_value=fake_response):
        result = get_parcel_status("OLD-FORMAT-123")

    assert result["state"] == "delivered"
    assert result["state_id"] == "xyz"


def test_refresh_tracking_saves_real_zr_dict_response_without_crashing(
    logged_in_admin, app, seed, monkeypatch
):
    """
    اختبار شامل من الطرف للطرف بنفس البيانات الحقيقية من production
    logs — الروت كامل، مو الدالة وحدها.
    """
    from unittest.mock import patch, MagicMock

    monkeypatch.setenv("ZR_TENANT_ID", "test-tenant")
    monkeypatch.setenv("ZR_API_KEY", "test-key")

    with app.app_context():
        order = Order(
            order_number="BD-REALTRK", customer_name="زبون حقيقي",
            customer_phone="0555000002", wilaya_code="24",
            wilaya_name="قالمة", subtotal=720, shipping_price=700,
            total=1420, status=Order.STATUS_CONFIRMED,
            tracking_number="FCQIKGY725-ZR-24",
        )
        db.session.add(order)
        db.session.commit()
        oid = order.id

    fake_response = MagicMock()
    fake_response.status_code = 200
    fake_response.content = b"{}"
    fake_response.headers = {}
    fake_response.json.return_value = {
        "state": {
            "id": "f34dfc2c-ecfe-4e06-9a19-430d2d0b27b9",
            "name": "encaisse",
            "description": "Encaissé",
            "isBlocking": False,
            "isLocked": False,
            "visibleFor": 15,
            "editableBy": 2,
            "color": "00796b",
        }
    }

    with patch("services.zrexpress.requests.request", return_value=fake_response):
        resp = logged_in_admin.post(
            f"/admin/orders/{oid}/refresh-tracking", follow_redirects=True
        )

    assert resp.status_code == 200

    with app.app_context():
        updated = db.session.get(Order, oid)
        assert updated.carrier_state == "Encaissé"
        assert isinstance(updated.carrier_state, str)


# ══════════════════════════════════════════════════════════════
# startup._fix_null_flags: قيم BOOLEAN لازم TRUE/FALSE ماشي 0/1
# ══════════════════════════════════════════════════════════════

def test_fix_null_flags_uses_ansi_boolean_literals(app):
    """
    بلاغ إنتاج حقيقي (v148): كنا نستعملو `0`/`1` كقيم افتراضية فـSQL
    خام لتصليح أعمدة BOOLEAN اللي بقات NULL. Postgres يرفض هذا:

        DatatypeMismatch: column "is_extra" is of type boolean
        but expression is of type integer

    SQLite ساهل ويقبل 0/1 كـBOOLEAN، فالمشكل ما كانش يبان غير
    فالإنتاج (Postgres) — وكان يتكرر فكل نشر بلا ما يتصلّح أبدًا.
    """
    from sqlalchemy import text
    from startup import _fix_null_flags
    from models import Category, ProductImage, ProductVariant

    with app.app_context():
        cat = Category(name="اختبار قيم", slug="test-boolean-flags")
        db.session.add(cat)
        db.session.flush()

        product = Product(name="منتج", slug="test-boolean-flags-p",
                          price=100, category_id=cat.id, is_active=True)
        db.session.add(product)
        db.session.flush()

        db.session.add(ProductImage(product_id=product.id, image_filename="x.jpg"))
        db.session.commit()

        # ⚠️ العمود الطازج عندو NOT NULL DEFAULT FALSE، فما نقدروش
        # نحطّو NULL مباشرة. نحاكيو الحالة الواقعية (قاعدة بيانات
        # قديمة قبل ما تنزاد القيود) بإعادة بناء العمود بلا قيد.
        db.session.execute(text(
            "ALTER TABLE product_images RENAME COLUMN is_extra TO is_extra_bak"
        ))
        db.session.execute(text(
            "ALTER TABLE product_images ADD COLUMN is_extra BOOLEAN"
        ))
        db.session.commit()

        # ما يطيحش، ويصلّح القيمة فعليًا
        _fix_null_flags(db)

        fixed = db.session.execute(text(
            "SELECT is_extra FROM product_images WHERE product_id = :pid"
        ), {"pid": product.id}).scalar()

        assert fixed in (0, False), "لازم يتصلّح لـFalse، ماشي يبقى NULL"


def test_fix_null_flags_source_uses_ansi_literals_not_integers():
    """
    ⚠️ SQLite متساهل مع 0/1 كـBOOLEAN، فاختبار السلوك وحدو ما يكفيش
    يكشف عطب Postgres-specific. نتحقق مباشرة من القيم النصية
    المستعملة فـSQL — هذا هو الضمان الحقيقي ضد رجوع الخطأ:
        DatatypeMismatch: column "is_extra" is of type boolean
        but expression is of type integer
    """
    import inspect
    from startup import _fix_null_flags

    source = inspect.getsource(_fix_null_flags)

    assert '"FALSE"' in source or "'FALSE'" in source
    assert '"TRUE"' in source or "'TRUE'" in source
    # ما نلقاوش "0" ولا "1" كقيم افتراضية فداخل قائمة targets
    assert '"0"),' not in source, "لازم TRUE/FALSE ماشي 0/1 (Postgres يرفضهم)"
    assert '"1"),' not in source, "لازم TRUE/FALSE ماشي 0/1 (Postgres يرفضهم)"


# ══════════════════════════════════════════════════════════════
# create_app: لازم تخدم مهما كان شكل رابط Postgres اللي يبعثو Railway
# ══════════════════════════════════════════════════════════════

def test_create_app_works_with_psycopg3_style_database_url(monkeypatch):
    """
    بلاغ إنتاج حرج (v149): الموقع طاح بالكامل —
        ModuleNotFoundError: No module named 'psycopg'

    السبب: DATABASE_URL عند Railway صار يبدأ بـ"postgresql+psycopg://"
    (سائق النسخة 3)، وكنا مثبّتين psycopg2-binary برك. كل عامل
    (worker) كان يطيح فور الإقلاع — outage كامل، ماشي عطب جزئي.

    هذا الاختبار يتأكد أن create_app() تقدر توصل لخطوة استيراد
    السائق بلا ما تطيح، بغض النظر عن شكل الرابط.
    """
    import importlib
    import secrets

    monkeypatch.setenv("SECRET_KEY", secrets.token_hex(32))
    monkeypatch.setenv(
        "DATABASE_URL", "postgresql+psycopg://user:pass@localhost:5432/fakedb"
    )

    import config
    importlib.reload(config)
    import app as app_module
    importlib.reload(app_module)

    # ما نطلبش اتصال حقيقي بقاعدة بيانات — نتأكد برك أن create_app()
    # توصل لآخرها بلا ModuleNotFoundError عند استيراد السائق (هنا
    # بالضبط كانت تطيح فالإنتاج، عند db.init_app → create_engine).
    created_app = app_module.create_app()
    assert created_app is not None


def test_create_app_works_with_psycopg2_style_database_url(monkeypatch):
    """الشكل القديم (بلا +psycopg صراحة) لازم يبقى يخدم، ماشي غير الجديد."""
    import importlib
    import secrets

    monkeypatch.setenv("SECRET_KEY", secrets.token_hex(32))
    monkeypatch.setenv(
        "DATABASE_URL", "postgresql://user:pass@localhost:5432/fakedb"
    )

    import config
    importlib.reload(config)
    import app as app_module
    importlib.reload(app_module)

    created_app = app_module.create_app()
    assert created_app is not None


# ══════════════════════════════════════════════════════════════
# v150: تحديث حالة الطلبية تلقائيًا من ZR + تعديل نجوم التقييم
# ══════════════════════════════════════════════════════════════

def test_map_zr_state_to_order_status():
    from services.zrexpress import map_state_to_order_status as m

    # «Encaissé» = ZR قبضات الكاش → مسلّمة (حالة حقيقية من الإنتاج)
    assert m("Encaissé") == "delivered"
    assert m("Livré") == "delivered"
    assert m("Recouvert") == "delivered"   # ZR قبضات المبلغ (من الإنتاج)
    # «Récupéré par fournisseur» = رجعت لصاحب المتجر (تأكيد منو)
    assert m("Récupéré par fournisseur") == "returned"
    # «Récupéré» وحدها (مثلاً من السائق) ماشي راجعة
    assert m("Récupéré par le livreur") != "returned"
    assert m("Retourné") == "returned"
    # «retour en cours» مازال ما رجعش
    assert m("Retour en cours") is None
    # ما نخلطوش «non livré» بـ«livré»
    assert m("Non livré") is None
    assert m("En attente") is None
    assert m("Expédié") == "shipped"
    assert m("En transit") == "shipped"
    # مكتب: يتحوّل لـ at_office غير للطلبيات نوعها مكتب
    assert m("Disponible au bureau", "office") == "at_office"
    assert m("Disponible au bureau", "home") == "shipped"
    # حالة مجهولة → ما نبدّلوش
    assert m("حاجة غريبة") is None
    assert m("") is None


def _zr_patch(label):
    from unittest.mock import patch
    return patch(
        "services.zrexpress.get_parcel_status",
        return_value={"state": label, "state_id": "x", "updated_at": None, "raw": {}},
    )


def _make_order(app, number, status, delivery_type="home"):
    with app.app_context():
        order = Order(
            order_number=number, customer_name="زبون", customer_phone="0555111222",
            wilaya_code="16", wilaya_name="الجزائر", subtotal=500,
            shipping_price=0, total=500, status=status,
            delivery_type=delivery_type, tracking_number="TRK-" + number,
        )
        db.session.add(order)
        db.session.commit()
        return order.id


def test_refresh_tracking_updates_order_status_forward(logged_in_admin, app, seed):
    oid = _make_order(app, "BD-V150A", Order.STATUS_CONFIRMED)
    with _zr_patch("Encaissé"):
        logged_in_admin.post(f"/admin/orders/{oid}/refresh-tracking",
                             follow_redirects=True)
    with app.app_context():
        o = db.session.get(Order, oid)
        assert o.carrier_state == "Encaissé"
        assert o.status == Order.STATUS_DELIVERED


def test_refresh_tracking_never_moves_status_backwards(logged_in_admin, app, seed):
    oid = _make_order(app, "BD-V150B", Order.STATUS_AT_OFFICE, "office")
    with _zr_patch("En transit"):   # shipped < at_office
        logged_in_admin.post(f"/admin/orders/{oid}/refresh-tracking",
                             follow_redirects=True)
    with app.app_context():
        assert db.session.get(Order, oid).status == Order.STATUS_AT_OFFICE


def test_refresh_tracking_leaves_delivered_and_cancelled_alone(
        logged_in_admin, app, seed):
    for number, st in (("BD-V150C", Order.STATUS_DELIVERED),
                       ("BD-V150D", Order.STATUS_CANCELLED)):
        oid = _make_order(app, number, st)
        with _zr_patch("Retourné"):
            logged_in_admin.post(f"/admin/orders/{oid}/refresh-tracking",
                                 follow_redirects=True)
        with app.app_context():
            assert db.session.get(Order, oid).status == st


def test_refresh_tracking_unknown_state_keeps_status(logged_in_admin, app, seed):
    oid = _make_order(app, "BD-V150E", Order.STATUS_PREPARED)
    with _zr_patch("حالة جديدة ما نعرفوهاش"):
        logged_in_admin.post(f"/admin/orders/{oid}/refresh-tracking",
                             follow_redirects=True)
    with app.app_context():
        o = db.session.get(Order, oid)
        assert o.status == Order.STATUS_PREPARED
        assert o.carrier_state == "حالة جديدة ما نعرفوهاش"


def test_admin_can_edit_review_rating(logged_in_admin, app, seed):
    from models import Review, AdminAction
    with app.app_context():
        r = Review(product_id=None, author_name="زبون", rating=1,
                   comment="ماشاء الله", is_approved=True)
        db.session.add(r)
        db.session.commit()
        rid = r.id

    resp = logged_in_admin.post(f"/admin/reviews/{rid}/rating",
                                data={"rating": "5"}, follow_redirects=True)
    assert resp.status_code == 200
    with app.app_context():
        assert db.session.get(Review, rid).rating == 5
        assert AdminAction.query.filter_by(action=AdminAction.REVIEW_EDIT).count() == 1

    # قيمة خارج المجال ما تتقبلش
    logged_in_admin.post(f"/admin/reviews/{rid}/rating", data={"rating": "9"})
    with app.app_context():
        assert db.session.get(Review, rid).rating == 5


def test_public_review_forms_have_no_default_star(client, app, seed):
    """الفورم ما لازمش يجي بنجوم مختارة مسبقًا — الزبون يختار بنفسو."""
    html = client.get("/reviews").get_data(as_text=True)
    assert 'name="rating"' in html and "checked" not in html.split('name="rating"', 1)[1].split("</form>")[0]


# ══════════════════════════════════════════════════════════════
# v151: زر «تحديث كل الطلبيات الممكنة»
# ══════════════════════════════════════════════════════════════

def test_bulk_refresh_updates_only_eligible_orders(logged_in_admin, app, seed, monkeypatch):
    from unittest.mock import patch
    monkeypatch.setenv("ZR_TENANT_ID", "t")
    monkeypatch.setenv("ZR_API_KEY", "k")

    a = _make_order(app, "BD-BLK-A", Order.STATUS_SHIPPED)       # لازم يتحدّث
    b = _make_order(app, "BD-BLK-B", Order.STATUS_DELIVERED)     # مسلّمة — نتجاهلوها
    c = _make_order(app, "BD-BLK-C", Order.STATUS_PREPARED)      # حالة مجهولة
    with app.app_context():
        d = Order(order_number="BD-BLK-D", customer_name="x", customer_phone="0555",
                  wilaya_code="16", wilaya_name="الجزائر", subtotal=1,
                  shipping_price=0, total=1, status=Order.STATUS_CONFIRMED)
        db.session.add(d); db.session.commit()
        did = d.id    # بلا رقم تتبّع — نتجاهلوها

    def fake(tracking):
        return {"state": "Livré" if tracking.endswith("BD-BLK-A") else "حاجة غريبة",
                "state_id": None, "updated_at": None, "raw": {}}

    with patch("services.zrexpress.get_parcel_status", side_effect=fake) as m:
        resp = logged_in_admin.post("/admin/orders/refresh-all-tracking",
                                    follow_redirects=True)
    assert resp.status_code == 200
    assert m.call_count == 2   # A و C برك

    with app.app_context():
        assert db.session.get(Order, a).status == Order.STATUS_DELIVERED
        assert db.session.get(Order, b).status == Order.STATUS_DELIVERED
        assert db.session.get(Order, c).status == Order.STATUS_PREPARED
        assert db.session.get(Order, c).carrier_state == "حاجة غريبة"
        assert db.session.get(Order, did).status == Order.STATUS_CONFIRMED


def test_bulk_refresh_one_failure_does_not_stop_others(logged_in_admin, app, seed, monkeypatch):
    from unittest.mock import patch
    from services.zrexpress import ZRError
    monkeypatch.setenv("ZR_TENANT_ID", "t")
    monkeypatch.setenv("ZR_API_KEY", "k")

    bad = _make_order(app, "BD-BLK-E", Order.STATUS_SHIPPED)
    good = _make_order(app, "BD-BLK-F", Order.STATUS_SHIPPED)

    def fake(tracking):
        if tracking.endswith("BD-BLK-E"):
            raise ZRError("timeout")
        return {"state": "Encaissé", "state_id": None, "updated_at": None, "raw": {}}

    with patch("services.zrexpress.get_parcel_status", side_effect=fake):
        resp = logged_in_admin.post("/admin/orders/refresh-all-tracking",
                                    follow_redirects=True)
    assert resp.status_code == 200
    assert "فشلت" in resp.get_data(as_text=True)
    with app.app_context():
        assert db.session.get(Order, bad).status == Order.STATUS_SHIPPED
        assert db.session.get(Order, good).status == Order.STATUS_DELIVERED


def test_bulk_refresh_respects_max_orders_cap(logged_in_admin, app, seed, monkeypatch):
    from unittest.mock import patch
    import routes.admin as admin_mod
    monkeypatch.setenv("ZR_TENANT_ID", "t")
    monkeypatch.setenv("ZR_API_KEY", "k")
    monkeypatch.setattr(admin_mod, "BULK_ZR_MAX_ORDERS", 2)

    for i in range(5):
        _make_order(app, f"BD-CAP-{i}", Order.STATUS_SHIPPED)

    with patch("services.zrexpress.get_parcel_status",
               return_value={"state": "En transit", "state_id": None,
                             "updated_at": None, "raw": {}}) as m:
        resp = logged_in_admin.post("/admin/orders/refresh-all-tracking",
                                    follow_redirects=True)
    assert m.call_count == 2
    assert "بقاو 3 طلبية" in resp.get_data(as_text=True)


def test_orders_list_shows_bulk_refresh_button(logged_in_admin, app, seed):
    html = logged_in_admin.get("/admin/orders").get_data(as_text=True)
    assert "تحديث كل الطلبيات الممكنة" in html


def test_bulk_refresh_reports_order_and_carrier_for_unknown_state(
        logged_in_admin, app, seed, monkeypatch):
    from unittest.mock import patch
    monkeypatch.setenv("ZR_TENANT_ID", "t")
    monkeypatch.setenv("ZR_API_KEY", "k")

    oid = _make_order(app, "BD-WRL-A", Order.STATUS_CONFIRMED)
    with app.app_context():
        o = db.session.get(Order, oid)
        o.carrier_code = "worldexpress"
        o.carrier_name = "World Express"
        db.session.commit()

    with patch("services.zrexpress.get_parcel_status",
               return_value={"state": "حالة غريبة", "state_id": None,
                             "updated_at": None, "raw": {}}):
        resp = logged_in_admin.post("/admin/orders/refresh-all-tracking",
                                    follow_redirects=True)
    page = resp.get_data(as_text=True)
    assert "BD-WRL-A" in page and "World Express" in page and "حالة غريبة" in page
    with app.app_context():
        assert db.session.get(Order, oid).status == Order.STATUS_CONFIRMED


# ══════════════════════════════════════════════════════════════
# v155: الزوار فصفحة وحدها
# ══════════════════════════════════════════════════════════════

def test_visitors_page_separate_from_dashboard(logged_in_admin, app, seed):
    dash = logged_in_admin.get("/admin/").get_data(as_text=True)
    assert "vis-panel" not in dash          # الزوار ما بقاوش فنظرة عامة
    assert 'href="/admin/visitors"' in dash  # بصح فيه رابط فالقائمة

    page = logged_in_admin.get("/admin/visitors")
    assert page.status_code == 200
    assert "الزوار" in page.get_data(as_text=True)


def test_visitors_reset_redirects_to_visitors_page(logged_in_admin, app, seed):
    resp = logged_in_admin.post("/admin/visitors/reset", data={"mode": "all"})
    assert resp.status_code == 302
    assert resp.headers["Location"].endswith("/admin/visitors")


def test_visitors_page_requires_login(client, app, seed):
    assert client.get("/admin/visitors").status_code in (302, 401)


# ══════════════════════════════════════════════════════════════
# v156: كوبون مكافأة شخصي للزبون
# ══════════════════════════════════════════════════════════════

def _reward_post(client, oid, percent="15", days="30"):
    data = {"percent": percent}
    if days is not None:
        data["days"] = days
    return client.post(f"/admin/orders/{oid}/reward-coupon", data=data,
                       follow_redirects=True)


def test_reward_coupon_created_for_delivered_order(logged_in_admin, app, seed):
    from models import Coupon, AdminAction
    oid = _make_order(app, "BD-RWD-A", Order.STATUS_DELIVERED)
    with app.app_context():
        db.session.get(Order, oid).customer_phone = "+213 555 11 12 22"
        db.session.commit()

    resp = _reward_post(logged_in_admin, oid, "15", "30")
    assert resp.status_code == 200

    with app.app_context():
        c = Coupon.query.filter_by(source_order_id=oid).one()
        import re
        assert re.fullmatch(r"[A-Z]{2}-[2-9A-HJ-NP-Z]{5}", c.code), c.code
        assert c.code.startswith("ZB-")        # اسم الطلبية «زبون» → ZB
        assert c.discount_type == Coupon.TYPE_PERCENT and c.value == 15
        assert c.usage_limit == 1                   # مرة وحدة
        assert c.customer_phone == "0555111222"     # صيغة موحّدة
        assert c.ends_at is not None
        assert AdminAction.query.filter_by(action=AdminAction.REWARD_COUPON).count() == 1


def test_reward_coupon_refused_for_undelivered_order(logged_in_admin, app, seed):
    from models import Coupon
    oid = _make_order(app, "BD-RWD-B", Order.STATUS_SHIPPED)
    _reward_post(logged_in_admin, oid)
    with app.app_context():
        assert Coupon.query.filter_by(source_order_id=oid).count() == 0


def test_reward_coupon_rejects_bad_percent_and_days(logged_in_admin, app, seed):
    from models import Coupon
    oid = _make_order(app, "BD-RWD-C", Order.STATUS_DELIVERED)
    for pct in ("0", "101", "-5", ""):
        _reward_post(logged_in_admin, oid, pct)
    _reward_post(logged_in_admin, oid, "10", "9999")
    with app.app_context():
        assert Coupon.query.filter_by(source_order_id=oid).count() == 0


def test_reward_coupon_not_duplicated_and_no_expiry_when_days_blank(
        logged_in_admin, app, seed):
    from models import Coupon
    oid = _make_order(app, "BD-RWD-D", Order.STATUS_DELIVERED)
    _reward_post(logged_in_admin, oid, "10", "")
    _reward_post(logged_in_admin, oid, "20", "")   # ثاني مرة — ما يتكرّرش
    with app.app_context():
        coupons = Coupon.query.filter_by(source_order_id=oid).all()
        assert len(coupons) == 1
        assert coupons[0].value == 10 and coupons[0].ends_at is None


def test_personal_coupon_only_valid_for_its_customer(app, seed):
    from models import Coupon
    with app.app_context():
        # ⚠️ نسجّلوه فالقاعدة باش الافتراضيات (min_order_amount...) تتعمّر
        c = Coupon(code="PERS-1", discount_type=Coupon.TYPE_PERCENT, value=10,
                   usage_limit=1, customer_phone="0555111222")
        normal = Coupon(code="NORM-1", discount_type=Coupon.TYPE_PERCENT, value=10)
        db.session.add_all([c, normal]); db.session.commit()

        # نفس الزبون بأي صيغة
        assert c.check_validity(1000, phone="0555111222")[0]
        assert c.check_validity(1000, phone="+213555111222")[0]
        assert c.check_validity(1000, phone="0555 11 12 22")[0]
        # زبون آخر
        ok, reason = c.check_validity(1000, phone="0666000000")
        assert not ok and "تاع زبون آخر" in reason
        # الرقم مازال مجهول (فالسلة) — ما نرفضوش
        assert c.check_validity(1000)[0]
        # كوبون عادي ما يتأثرش
        assert normal.check_validity(1000, phone="0666000000")[0]


def test_create_order_enforces_personal_coupon_phone(app, seed):
    import pytest
    from models import Coupon
    from services.orders import create_order, OrderError
    from services.cart import CartLine
    with app.app_context():
        c = Coupon(code="PERS-2", discount_type=Coupon.TYPE_PERCENT, value=10,
                   usage_limit=1, customer_phone="0555111222")
        db.session.add(c); db.session.commit()
        line = CartLine(key="product:1:0", kind="product", name="x",
                        unit_price=100, quantity=1, product_id=1)
        with pytest.raises(OrderError) as e:
            create_order(
                lines=[line],
                customer_data={"name": "غريب", "phone": "0666000000",
                               "wilaya_code": "39", "wilaya_name": "الوادي",
                               "commune": "الوادي", "address_note": None},
                shipping_info={"carrier_code": "zrexpress", "carrier_name": "ZR",
                               "delivery_type": "home", "price": 600},
                coupon=Coupon.query.filter_by(code="PERS-2").first(),
            )
        assert "تاع زبون آخر" in str(e.value)


def test_reward_message_and_button(logged_in_admin, app, seed):
    from models import Coupon
    oid = _make_order(app, "BD-RWD-E", Order.STATUS_DELIVERED)
    _reward_post(logged_in_admin, oid, "12", "14")
    with app.app_context():
        order = db.session.get(Order, oid)
        c = order.reward_coupon
        msg = order.reward_coupon_message(c, "Budac", "https://budacdz.com")
        assert c.code in msg and "12%" in msg and order.order_number in msg
        assert "مرة وحدة" in msg
        # ⚠️ لهجة جزائرية: ما نستعملوش كلمات مغربية
        words = set(msg.replace("\n", " ").replace("*", " ").split())
        for banned in ("ديال", "دابا", "بغيت", "غادي", "شنو", "مزيان", "فين",
                       "غير", "خاص", "مش", "عجباتك", "وصلاتك", "وصلات", "تحوّزها"):
            assert banned not in words

    html = logged_in_admin.get(f"/admin/orders/{oid}").get_data(as_text=True)
    assert "ابعث الكوبون عبر واتساب" in html and "wa.me/" in html


def test_reward_panel_hidden_for_undelivered(logged_in_admin, app, seed):
    oid = _make_order(app, "BD-RWD-F", Order.STATUS_CONFIRMED)
    html = logged_in_admin.get(f"/admin/orders/{oid}").get_data(as_text=True)
    assert "مكافأة الزبون" not in html


def test_latin_initials_for_codes():
    from utils import latin_initials as li
    assert li("Rania allem") == "RA"
    assert li("Youssra Salah salah") == "YO"
    assert li("Éric") == "ER"
    assert li("كمال") == "KM"
    assert li("ميساء ص") == "MY"
    assert li("شيماء") == "SH"
    assert li("عبادلية") == "AB"
    assert li("أستغفر الله") == "AS"
    # ما فيهش حروف يتحوّلو → BD
    assert li("") == "BD" and li(None) == "BD" and li("123 !!") == "BD"
    # حرف وحد → نكمّلو بالاحتياطي
    assert len(li("A")) == 2 and li("A").startswith("A")


def test_reward_code_uses_customer_initials(logged_in_admin, app, seed):
    from models import Coupon
    oid = _make_order(app, "BD-RWD-G", Order.STATUS_DELIVERED)
    with app.app_context():
        db.session.get(Order, oid).customer_name = "Rania allem"
        db.session.commit()
    _reward_post(logged_in_admin, oid, "10", "30")
    with app.app_context():
        assert Coupon.query.filter_by(source_order_id=oid).one().code.startswith("RA-")


# ══════════════════════════════════════════════════════════════
# v158: كوبونات مكافأة جماعية من صفحة الطلبيات
# ══════════════════════════════════════════════════════════════

def _make_order_for(app, number, status, name, phone):
    oid = _make_order(app, number, status)
    with app.app_context():
        o = db.session.get(Order, oid)
        o.customer_name, o.customer_phone = name, phone
        db.session.commit()
    return oid


def test_bulk_reward_creates_one_coupon_per_customer(logged_in_admin, app, seed):
    from models import Coupon
    a = _make_order_for(app, "BD-BR-A", Order.STATUS_DELIVERED, "Rania allem", "0556329970")
    b = _make_order_for(app, "BD-BR-B", Order.STATUS_DELIVERED, "Rania allem", "+213556329970")  # نفس الزبون
    c = _make_order_for(app, "BD-BR-C", Order.STATUS_DELIVERED, "كمال مجذوب", "0662834652")
    d = _make_order_for(app, "BD-BR-D", Order.STATUS_SHIPPED, "Diaf Akila", "0777000111")        # ماشي مسلّمة

    resp = logged_in_admin.post(
        "/admin/orders/bulk-reward",
        data={"order_ids": [a, b, c, d], "percent": "10", "days": "30"})
    assert resp.status_code == 302
    assert "/admin/orders/reward-batch?ids=" in resp.headers["Location"]

    with app.app_context():
        coupons = Coupon.query.filter(Coupon.source_order_id.in_([a, b, c, d])).all()
        by_order = {cp.source_order_id: cp for cp in coupons}
        assert set(by_order) == {a, c}               # b نفس زبون a — d ماشي مسلّمة
        assert by_order[a].code.startswith("RA-")
        assert by_order[c].code.startswith("KM-")
        assert all(cp.value == 10 and cp.usage_limit == 1 for cp in coupons)

    page = logged_in_admin.get(resp.headers["Location"], follow_redirects=True)
    html = page.get_data(as_text=True)
    assert html.count("wa.me/") == 2
    assert "BD-BR-A" in html and "BD-BR-C" in html and "BD-BR-D" not in html


def test_bulk_reward_validates_percent_and_selection(logged_in_admin, app, seed):
    from models import Coupon
    a = _make_order_for(app, "BD-BR-E", Order.STATUS_DELIVERED, "Test", "0555000111")
    r = logged_in_admin.post("/admin/orders/bulk-reward",
                             data={"order_ids": [a], "percent": "0"})
    assert r.status_code == 302 and "reward-batch" not in r.headers["Location"]
    r = logged_in_admin.post("/admin/orders/bulk-reward", data={"percent": "10"})
    assert r.status_code == 302 and "reward-batch" not in r.headers["Location"]
    with app.app_context():
        assert Coupon.query.filter_by(source_order_id=a).count() == 0


def test_bulk_reward_reuses_existing_active_coupon(logged_in_admin, app, seed):
    from models import Coupon
    a = _make_order_for(app, "BD-BR-F", Order.STATUS_DELIVERED, "Samir", "0555123123")
    logged_in_admin.post("/admin/orders/bulk-reward", data={"order_ids": [a], "percent": "10"})
    logged_in_admin.post("/admin/orders/bulk-reward", data={"order_ids": [a], "percent": "20"})
    with app.app_context():
        cps = Coupon.query.filter_by(source_order_id=a).all()
        assert len(cps) == 1 and cps[0].value == 10    # ما يتكرّرش ولا يتبدّل


def test_bulk_reward_all_undelivered_creates_nothing(logged_in_admin, app, seed):
    from models import Coupon
    a = _make_order_for(app, "BD-BR-G", Order.STATUS_CONFIRMED, "Yacine", "0555321321")
    r = logged_in_admin.post("/admin/orders/bulk-reward", data={"order_ids": [a], "percent": "10"})
    assert "reward-batch" not in r.headers["Location"]
    with app.app_context():
        assert Coupon.query.filter_by(source_order_id=a).count() == 0


def test_orders_list_has_bulk_reward_button(logged_in_admin, app, seed):
    _make_order_for(app, "BD-BR-H", Order.STATUS_DELIVERED, "Test", "0555000999")
    html = logged_in_admin.get("/admin/orders").get_data(as_text=True)
    assert "كوبون مكافأة للمختارة" in html and "bulk-reward" in html


def test_reward_batch_requires_login(client, app, seed):
    assert client.get("/admin/orders/reward-batch?ids=1").status_code in (302, 401)
    assert client.post("/admin/orders/bulk-reward", data={"order_ids": [1], "percent": "10"}).status_code in (302, 401, 400)


def test_bulk_reward_handles_duplicated_checkbox_ids(logged_in_admin, app, seed):
    """الجدول والبطاقة يبعثو نفس id مرتين — لازم يتعالج كطلبية وحدة."""
    from models import Coupon
    a = _make_order_for(app, "BD-DUP-A", Order.STATUS_DELIVERED, "Samia", "0555777888")
    r = logged_in_admin.post("/admin/orders/bulk-reward",
                             data={"order_ids": [a, a, a, a], "percent": "10"})
    assert f"ids={a}" in r.headers["Location"]
    with app.app_context():
        assert Coupon.query.filter_by(source_order_id=a).count() == 1


def test_orders_list_counts_unique_orders_in_js(logged_in_admin, app, seed):
    _make_order_for(app, "BD-DUP-B", Order.STATUS_DELIVERED, "Test", "0555000222")
    html = logged_in_admin.get("/admin/orders").get_data(as_text=True)
    assert "uniqueValues" in html     # العدّ بالقيم المختلفة، ماشي بعدد الصناديق


# ══════════════════════════════════════════════════════════════
# v160: ضغط gzip + كاش طويل للملفات الثابتة
# ══════════════════════════════════════════════════════════════

def test_public_pages_are_gzipped_when_browser_accepts(client, app, seed):
    import gzip
    r = client.get("/", headers={"Accept-Encoding": "gzip"})
    assert r.status_code == 200
    assert r.headers.get("Content-Encoding") == "gzip"
    assert "Accept-Encoding" in r.headers.get("Vary", "")
    assert b"<html" in gzip.decompress(r.data)          # المحتوى سليم بعد فك الضغط


def test_no_gzip_when_browser_does_not_accept(client, app, seed):
    r = client.get("/", headers={"Accept-Encoding": "identity"})
    assert r.status_code == 200
    assert "Content-Encoding" not in r.headers


def test_admin_pages_are_never_compressed(client, app, seed):
    r = client.get("/admin/login", headers={"Accept-Encoding": "gzip"})
    assert r.status_code == 200
    assert "Content-Encoding" not in r.headers


def test_static_css_is_gzipped(client, app, seed):
    import gzip
    r = client.get("/static/css/style.css", headers={"Accept-Encoding": "gzip"})
    assert r.status_code == 200 and r.headers.get("Content-Encoding") == "gzip"
    original = open("static/css/style.css", "rb").read()
    assert gzip.decompress(r.data) == original            # نفس الملف بالضبط
    assert len(r.data) < len(original) * 0.5               # صغر بنسبة معتبرة


def test_fingerprinted_static_gets_long_immutable_cache(client, app, seed):
    html = client.get("/").get_data(as_text=True)
    import re
    m = re.search(r'/static/css/style\.css\?v=(\d+)', html)
    assert m, "رابط CSS لازم يجي ببصمة"
    r = client.get(f"/static/css/style.css?v={m.group(1)}")
    cc = r.headers.get("Cache-Control", "")
    assert "max-age=31536000" in cc and "immutable" in cc and "public" in cc


def test_unversioned_css_is_not_cached_forever(client, app, seed):
    r = client.get("/static/css/style.css")
    assert "immutable" not in r.headers.get("Cache-Control", "")


def test_static_images_without_fingerprint_get_30_days(client, app, seed):
    r = client.get("/static/img/brand/logo.png")
    assert r.status_code == 200
    assert "max-age=2592000" in r.headers.get("Cache-Control", "")


def test_no_double_version_param_in_rendered_html(client, app, seed):
    import re
    html = client.get("/").get_data(as_text=True)
    assert not re.search(r'\?v=[^"\'\s]*\?v=', html)       # كان «favicon.ico?v=123?v=2»


# ══════════════════════════════════════════════════════════════
# v161: نماذج رسائل واتساب قابلة للتعديل + حد أدنى لكوبون المكافأة
# ══════════════════════════════════════════════════════════════

def test_every_default_message_starts_with_brand_greeting(app, seed):
    """كل رسالة تبدا بـ«معاك BUDAC» (طلب صاحب المتجر)."""
    from services import wa_templates as wa
    with app.app_context():
        oid = _make_order(app, "BD-WA-A", Order.STATUS_DELIVERED)
        order = db.session.get(Order, oid)
        for spec in wa.KINDS:
            if spec["key"] in ("reward", "abandoned"):
                continue
            msg = order.whatsapp_message(spec["key"], "ignored-store-name")
            assert "معاك BUDAC." in msg, spec["key"]
            assert msg.startswith("سلام"), spec["key"]
            assert "{" not in msg, f"متغيّر ما تعوّضش فـ{spec['key']}: {msg}"


def test_reward_message_has_brand_and_min_amount(app, seed):
    from models import Coupon
    with app.app_context():
        oid = _make_order(app, "BD-WA-B", Order.STATUS_DELIVERED)
        order = db.session.get(Order, oid)
        c = Coupon(code="AM-TEST1", discount_type=Coupon.TYPE_PERCENT, value=10,
                   usage_limit=1, min_order_amount=5000)
        db.session.add(c); db.session.commit()
        msg = order.reward_coupon_message(c, base_url="https://budacdz.com")
        assert "معاك BUDAC." in msg and "AM-TEST1" in msg
        assert "5,000 دج" in msg and "{" not in msg
        c.min_order_amount = 0
        assert "فما فوق" not in order.reward_coupon_message(c)


def test_admin_edit_template_changes_the_sent_message(logged_in_admin, app, seed):
    from services import wa_templates as wa
    oid = _make_order(app, "BD-WA-C", Order.STATUS_SHIPPED)
    text = "سلام {name}، معاك {store}. طلبيتك {order_number} في الطريق {truck}"
    r = logged_in_admin.post("/admin/whatsapp-templates/shipped", data={"text": text},
                             follow_redirects=True)
    assert r.status_code == 200
    with app.app_context():
        assert wa.is_customized("shipped")
        msg = db.session.get(Order, oid).whatsapp_message("shipped")
        assert msg.startswith("سلام زبون، معاك BUDAC. طلبيتك BD-WA-C في الطريق")
    # الصفحة تبيّن النص الجديد + علامة «معدّلة»
    html = logged_in_admin.get("/admin/whatsapp-templates").get_data(as_text=True)
    assert "معدّلة" in html and "في الطريق" in html


def test_admin_reset_template_restores_default(logged_in_admin, app, seed):
    from services import wa_templates as wa
    logged_in_admin.post("/admin/whatsapp-templates/thanks", data={"text": "نص جديد {name}"})
    with app.app_context():
        assert wa.is_customized("thanks")
    logged_in_admin.post("/admin/whatsapp-templates/thanks/reset")
    with app.app_context():
        assert not wa.is_customized("thanks")
        assert wa.get_template("thanks") == wa.KIND_BY_KEY["thanks"]["default"]


def test_template_validation_rules(logged_in_admin, app, seed):
    from services import wa_templates as wa
    with app.app_context():
        v = wa.validate_template
        assert v("shipped", "") is not None                        # فارغ
        assert v("shipped", "x" * (wa.MAX_LEN + 1)) is not None    # طويل بزاف
        assert v("shipped", "مرحبا 😀") is not None                # إيموجي 4 بايت
        assert v("shipped", "تم ✅") is not None                    # إيموجي ملوّن BMP
        assert v("shipped", "تم {check}") is None                  # رمز نصّي مقبول
        assert v("reward", "خصم بلا كود") is not None             # {code} ضروري
        assert v("review", "قيّمنا") is not None                   # {link} ضروري
        assert v("reward", "الكود {code}") is None
        assert v("unknown_kind", "x") is not None
    # الرفض يوصل للأدمن ولا يحفظ شي
    logged_in_admin.post("/admin/whatsapp-templates/shipped", data={"text": "مرحبا 😀"})
    with app.app_context():
        assert not wa.is_customized("shipped")


def test_unknown_placeholders_never_break_sending(app, seed):
    from services import wa_templates as wa
    out = wa.render_text("سلام {name} {nothing_here} {", {"name": "علي"})
    assert out == "سلام علي {nothing_here} {"


def test_wa_templates_page_and_nav(logged_in_admin, app, seed):
    html = logged_in_admin.get("/admin/whatsapp-templates").get_data(as_text=True)
    for label in ("تأكيد الطلبية", "الطلبية خرجت", "كوبون مكافأة", "سلة متروكة"):
        assert label in html
    assert "معاك BUDAC" in html
    dash = logged_in_admin.get("/admin/").get_data(as_text=True)
    assert "رسائل واتساب" in dash


def test_wa_templates_require_login(client, app, seed):
    assert client.get("/admin/whatsapp-templates").status_code in (302, 401)
    assert client.post("/admin/whatsapp-templates/thanks", data={"text": "x"}).status_code in (302, 401, 400)


def test_reward_coupon_min_amount_single_and_bulk(logged_in_admin, app, seed):
    from models import Coupon
    a = _make_order_for(app, "BD-MIN-A", Order.STATUS_DELIVERED, "Sara", "0555444333")
    logged_in_admin.post(f"/admin/orders/{a}/reward-coupon",
                         data={"percent": "10", "days": "30", "min_amount": "5000"})
    b = _make_order_for(app, "BD-MIN-B", Order.STATUS_DELIVERED, "Hamza", "0555444222")
    c = _make_order_for(app, "BD-MIN-C", Order.STATUS_DELIVERED, "Nour", "0555444111")
    logged_in_admin.post("/admin/orders/bulk-reward",
                         data={"order_ids": [b, c], "percent": "15", "min_amount": "3000"})
    with app.app_context():
        assert Coupon.query.filter_by(source_order_id=a).one().min_order_amount == 5000
        for oid in (b, c):
            assert Coupon.query.filter_by(source_order_id=oid).one().min_order_amount == 3000
    # القيمة السالبة أو غير الرقمية ترفض
    d = _make_order_for(app, "BD-MIN-D", Order.STATUS_DELIVERED, "Zaki", "0555444000")
    for bad in ("-5", "abc", "99999999"):
        logged_in_admin.post(f"/admin/orders/{d}/reward-coupon",
                             data={"percent": "10", "min_amount": bad})
    with app.app_context():
        assert Coupon.query.filter_by(source_order_id=d).count() == 0


def test_personal_coupon_with_min_amount_is_enforced_at_validity(app, seed):
    from models import Coupon
    with app.app_context():
        c = Coupon(code="MIN-1", discount_type=Coupon.TYPE_PERCENT, value=10,
                   usage_limit=1, min_order_amount=5000)
        db.session.add(c); db.session.commit()
        assert not c.check_validity(4000)[0]
        assert c.check_validity(5000)[0]


# ══════════════════════════════════════════════════════════════
# v162: محرّر الصور (قص وتدوير) + إصلاح CSP للصور من blob
# ══════════════════════════════════════════════════════════════

def test_csp_allows_blob_images(client, app, seed):
    """بلاه blob: المتصفح يمنع تصغير الصور ومعاينة القص بصمت."""
    csp = client.get("/").headers.get("Content-Security-Policy", "")
    img_src = [d for d in csp.split(";") if d.strip().startswith("img-src")][0]
    assert "blob:" in img_src and "'self'" in img_src
    assert "script-src 'self'" in csp                      # المكتبة محلية، ماشي CDN


def test_cropper_vendor_files_are_served(client, app, seed):
    for path, marker in (
        ("/static/vendor/cropperjs/cropper.min.js", b"Cropper.js v1.6.2"),
        ("/static/vendor/cropperjs/cropper.min.css", b".cropper-container"),
        ("/static/vendor/cropperjs/LICENSE", b"MIT"),
        ("/static/js/image-editor.js", b"BudacImageEditor"),
    ):
        r = client.get(path)
        assert r.status_code == 200, path
        assert marker in r.data, path


def test_admin_pages_load_image_editor_script(logged_in_admin, app, seed):
    html = logged_in_admin.get("/admin/").get_data(as_text=True)
    assert "js/image-editor.js" in html


def test_product_form_routes_uploads_through_editor(logged_in_admin, app, seed):
    html = logged_in_admin.get("/admin/products/new").get_data(as_text=True)
    assert "BudacImageEditor.edit" in html
    assert "aspect: box.dataset.kind === 'extra' ? NaN : 1" in html   # الأساسية 1:1، التوضيحية حرّة


def test_category_and_offer_file_inputs_use_editor(logged_in_admin, app, seed):
    cats = logged_in_admin.get("/admin/categories").get_data(as_text=True)
    assert 'name="image"' in cats and 'data-image-editor="1"' in cats
    offer = logged_in_admin.get("/admin/offers/new").get_data(as_text=True)
    assert 'name="banner"' in offer and 'data-image-editor="16/9"' in offer


# ══════════════════════════════════════════════════════════════
# v163: استيراد المنتجات مع صورها (Excel + ZIP)
# ══════════════════════════════════════════════════════════════

def _xl_bytes(products, variants=(), images=()):
    """Excel صغير بنفس صيغة الاستيراد: products=(slug,name,price)، variants=(slug,label)، images=(slug,label,file)."""
    import io
    from openpyxl import Workbook
    from services.excel_io import (COLS_PRODUCTS, COLS_VARIANTS, COLS_IMAGES,
                                   SHEET_PRODUCTS, SHEET_VARIANTS, SHEET_IMAGES)
    wb = Workbook()
    ws = wb.active; ws.title = SHEET_PRODUCTS
    ws.append(COLS_PRODUCTS)
    for slug, name, price in products:
        row = {c: None for c in COLS_PRODUCTS}
        row.update({"الرابط (slug)": slug, "اسم المنتج": name, "السعر": price, "مفعّل": "لا"})
        ws.append([row[c] for c in COLS_PRODUCTS])
    wv = wb.create_sheet(SHEET_VARIANTS); wv.append(COLS_VARIANTS)
    for slug, label in variants:
        row = {c: None for c in COLS_VARIANTS}
        row.update({"رابط المنتج": slug, "اسم الموديل": label})
        wv.append([row[c] for c in COLS_VARIANTS])
    wi = wb.create_sheet(SHEET_IMAGES); wi.append(COLS_IMAGES)
    for r in images:
        wi.append(list(r))
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _jpeg_bytes(color=(200, 30, 30), size=(300, 200)):
    import io
    from PIL import Image
    buf = io.BytesIO(); Image.new("RGB", size, color).save(buf, "JPEG"); return buf.getvalue()


def _zip_bytes(files):
    import io, zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name, data in files.items():
            z.writestr(name, data)
    return buf.getvalue()


def _post_import(client, xlsx, zip_bytes=None, zip_name="imgs.zip"):
    import io
    data = {"excel_file": (io.BytesIO(xlsx), "p.xlsx")}
    if zip_bytes is not None:
        data["images_zip"] = (io.BytesIO(zip_bytes), zip_name)
    return client.post("/admin/products/import", data=data,
                       content_type="multipart/form-data", follow_redirects=True)


def test_import_with_zip_attaches_images_and_links_variants(logged_in_admin, app, seed):
    from models import Product, ProductImage
    xlsx = _xl_bytes(
        [("zz-handle", "مقبض تجريبي", 500), ("zz-single", "منتج وحيد", 300)],
        variants=[("zz-handle", "ذهبي"), ("zz-handle", "فضي")],
        images=[("zz-handle", "ذهبي", "gold.jpg"), ("zz-handle", "فضي", "silver.jpg"),
                ("zz-single", "", "single.jpg")],
    )
    zb = _zip_bytes({"folder/gold.jpg": _jpeg_bytes((220, 180, 40)),   # مجلد داخل ZIP — عادي
                     "silver.jpg": _jpeg_bytes((190, 190, 190)),
                     "single.jpg": _jpeg_bytes((10, 10, 200)),
                     "unused.jpg": _jpeg_bytes()})                      # ماشي فـExcel ← ما يتلمسش
    r = _post_import(logged_in_admin, xlsx, zb)
    assert r.status_code == 200
    with app.app_context():
        handle = Product.query.filter_by(slug="zz-handle").one()
        single = Product.query.filter_by(slug="zz-single").one()
        assert ProductImage.query.filter_by(product_id=handle.id).count() == 2
        assert ProductImage.query.filter_by(product_id=single.id).count() == 1
        assert handle.image_filename and single.image_filename            # الرئيسية تتحدّد
        imgs = {v.label: v.image_filename for v in handle.variants}
        assert imgs["ذهبي"] and imgs["فضي"] and imgs["ذهبي"] != imgs["فضي"]
        assert ProductImage.query.count() == 3                             # unused.jpg ما دخلش


def test_import_images_is_idempotent_and_keeps_manual_variant_choice(logged_in_admin, app, seed):
    from models import Product, ProductImage, ProductVariant
    xlsx = _xl_bytes([("zz-idem", "منتج", 100)], [("zz-idem", "أحمر")],
                     [("zz-idem", "أحمر", "red.jpg")])
    zb = _zip_bytes({"red.jpg": _jpeg_bytes()})
    _post_import(logged_in_admin, xlsx, zb)
    with app.app_context():
        v = ProductVariant.query.filter_by(label="أحمر").one()
        v.image_filename = "choix-manuel.webp"           # الأدمن بدّل الصورة يدويًا
        db.session.commit()
    _post_import(logged_in_admin, xlsx, zb)               # نفس الملفين مرة ثانية
    with app.app_context():
        p = Product.query.filter_by(slug="zz-idem").one()
        assert ProductImage.query.filter_by(product_id=p.id).count() == 1   # ما تكرّرتش
        assert ProductVariant.query.filter_by(label="أحمر").one().image_filename == "choix-manuel.webp"


def test_import_reports_missing_unknown_and_bad_images(logged_in_admin, app, seed):
    from models import ProductImage
    xlsx = _xl_bytes([("zz-warn", "منتج", 100)], [],
                     [("zz-warn", "", "absent.jpg"),          # ماكانش فـZIP
                      ("no-such-product", "", "a.jpg"),       # منتج مجهول
                      ("zz-warn", "", "fake.jpg"),            # ماشي صورة حقيقية
                      ("zz-warn", "موديل مجهول", "ok.jpg")])  # صورة تتزاد + موديل مجهول
    zb = _zip_bytes({"a.jpg": _jpeg_bytes(), "fake.jpg": b"this is not an image",
                     "ok.jpg": _jpeg_bytes()})
    html = _post_import(logged_in_admin, xlsx, zb).get_data(as_text=True)
    assert "absent.jpg" in html and "no-such-product" in html
    assert "fake.jpg" in html and "موديل مجهول" in html
    with app.app_context():
        assert ProductImage.query.count() == 1               # ok.jpg برك


def test_import_rejects_rar_and_non_zip_but_keeps_excel_untouched(logged_in_admin, app, seed):
    from models import Product
    xlsx = _xl_bytes([("zz-rar", "منتج", 100)])
    html = _post_import(logged_in_admin, xlsx, b"Rar!\x1a\x07\x00junk", "photos.rar").get_data(as_text=True)
    assert "RAR" in html and "ZIP" in html
    html = _post_import(logged_in_admin, xlsx, b"not a zip at all", "photos.zip").get_data(as_text=True)
    assert "ماشي ZIP صحيح" in html
    with app.app_context():
        assert Product.query.filter_by(slug="zz-rar").count() == 0   # رفض كامل قبل أي كتابة


def test_import_warns_when_sheet_and_zip_do_not_match(logged_in_admin, app, seed):
    # صفحة «الصور» بلا ZIP
    xlsx = _xl_bytes([("zz-nozip", "منتج", 100)], [], [("zz-nozip", "", "a.jpg")])
    html = _post_import(logged_in_admin, xlsx).get_data(as_text=True)
    assert "بلا ملف ZIP" in html
    # ZIP بلا صفحة «الصور»
    xlsx2 = _xl_bytes([("zz-nosheet", "منتج", 100)])
    html = _post_import(logged_in_admin, xlsx2, _zip_bytes({"a.jpg": _jpeg_bytes()})).get_data(as_text=True)
    assert "بلا صفحة" in html


def test_zip_slip_names_never_touch_disk_and_names_are_repaired(app):
    import zipfile, io
    from services.excel_io import _zip_basename, _zip_index, _find_zip_member
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("../../etc/evil.jpg", b"x")
        z.writestr("__MACOSX/._a.jpg", b"x")
        z.writestr("Dir\\Photo One.JPG", b"x")
    with zipfile.ZipFile(buf) as z:
        idx = _zip_index(z)
        assert "evil.jpg" in idx and "photo one.jpg" in idx and "._a.jpg" not in idx
        assert _find_zip_member(idx, "photo one") is not None           # بلا امتداد + حروف كبيرة/صغيرة
        assert _find_zip_member(idx, "PHOTO ONE.jpg") is not None
    # اسم عربي مخزّن بـUTF-8 بلا علامة (ZIP قديم) يتصلّح
    info = zipfile.ZipInfo("مقبض ذهبي.jpg".encode("utf-8").decode("cp437"))
    info.flag_bits = 0
    assert _zip_basename(info) == "مقبض ذهبي.jpg"


def test_import_route_allows_bigger_request_than_global_limit(logged_in_admin, app, seed):
    """الحد العام 30 ميغا ما يوقّفش استيراد Excel + صور."""
    from models import Product
    app.config["MAX_CONTENT_LENGTH"] = 150_000           # حد عام صغير للتجربة
    xlsx = _xl_bytes([("zz-big", "منتج", 100)], [], [("zz-big", "", "big.jpg")])
    from PIL import Image
    import io, os
    noise = Image.frombytes("RGB", (400, 400), os.urandom(400 * 400 * 3))   # صورة ما تتضغطش
    b = io.BytesIO(); noise.save(b, "PNG")
    zb = _zip_bytes({"big.jpg": b.getvalue()})
    assert len(zb) > 150_000
    r = _post_import(logged_in_admin, xlsx, zb)
    assert r.status_code == 200                                         # ماشي 413
    with app.app_context():
        assert Product.query.filter_by(slug="zz-big").count() == 1


def test_products_page_has_zip_input(logged_in_admin, app, seed):
    html = logged_in_admin.get("/admin/products").get_data(as_text=True)
    assert 'name="images_zip"' in html and 'accept=".zip"' in html


# ══════════════════════════════════════════════════════════════
# v164: تفعيل دفعة وحدة + تعديل الصور الموجودة (قص/تدوير/قلب)
# ══════════════════════════════════════════════════════════════

def _mk_product(app, slug, price=100, active=False, image=None, name=None):
    from models import Product
    with app.app_context():
        p = Product(name=name or slug, slug=slug, price=price, is_active=active,
                    image_filename=image)
        db.session.add(p); db.session.commit()
        return p.id


def _bulk(client, action, ids, back=None):
    data = {"action": action, "product_ids": ids}
    if back is not None:
        data["back"] = back
    return client.post("/admin/products/bulk-status", data=data)


def test_bulk_activate_products_with_guards(logged_in_admin, app, seed):
    from models import Product, AdminAction
    a = _mk_product(app, "bk-a", 100, image="x.webp", name="مقبض أ")
    b = _mk_product(app, "bk-b", 100, image=None, name="مقبض ب")        # بلا صورة ← يتفعّل مع تحذير
    c = _mk_product(app, "bk-c", 0, image="y.webp", name="بلا سعر")      # السعر 0 ← ما يتفعّلش
    r = _bulk(logged_in_admin, "activate", [a, a, b, c, b])               # تكرار (checkbox مزدوجة)
    assert r.status_code == 302
    with app.app_context():
        assert db.session.get(Product, a).is_active is True
        assert db.session.get(Product, b).is_active is True
        assert db.session.get(Product, c).is_active is False
        assert AdminAction.query.filter_by(action=AdminAction.PRODUCT_BULK).count() == 1
    html = logged_in_admin.get("/admin/products").get_data(as_text=True)
    assert "تفعّل 2 منتج" in html and "بلا سعر" in html and "بلا صورة" in html


def test_bulk_deactivate_and_noop_cases(logged_in_admin, app, seed):
    from models import Product
    a = _mk_product(app, "bk-d", 100, active=True)
    b = _mk_product(app, "bk-e", 100, active=False)
    _bulk(logged_in_admin, "deactivate", [a, b])
    with app.app_context():
        assert db.session.get(Product, a).is_active is False
    # إجراء مجهول / بلا اختيار ← ما يتبدّل والو
    _bulk(logged_in_admin, "delete-everything", [a])
    _bulk(logged_in_admin, "activate", [])
    with app.app_context():
        assert db.session.get(Product, a).is_active is False


def test_bulk_status_back_redirect_is_safe(logged_in_admin, app, seed):
    a = _mk_product(app, "bk-f", 100)
    ok = _bulk(logged_in_admin, "activate", [a], back="/admin/products?status=inactive")
    assert ok.headers["Location"].endswith("/admin/products?status=inactive")
    evil = _bulk(logged_in_admin, "activate", [a], back="https://evil.example/steal")
    assert "evil.example" not in evil.headers["Location"]
    assert evil.headers["Location"].endswith("/admin/products")


def test_bulk_status_requires_login(client, app, seed):
    r = client.post("/admin/products/bulk-status", data={"action": "activate", "product_ids": [1]})
    assert r.status_code in (302, 401, 400)


def test_products_list_has_checkboxes_and_bulk_bar(logged_in_admin, app, seed):
    _mk_product(app, "bk-g", 100, active=False)
    html = logged_in_admin.get("/admin/products?status=inactive").get_data(as_text=True)
    assert 'id="bulkProductsForm"' in html and 'form="bulkProductsForm"' in html
    assert "تفعيل المختارة" in html and "تعطيل المختارة" in html and 'id="bulkAll"' in html


def _upload_via_route(client, pid, color):
    import io
    return client.post(f"/admin/products/{pid}/images/upload",
                       data={"image": (io.BytesIO(_jpeg_bytes(color)), "p.jpg"), "kind": "main"},
                       content_type="multipart/form-data").get_json()


def test_replace_product_image_keeps_position_main_and_variants(logged_in_admin, app, seed):
    import io, os
    from models import Product, ProductImage, ProductVariant
    from services.uploads import get_upload_dir
    pid = _mk_product(app, "rp-a", 100)
    first = _upload_via_route(logged_in_admin, pid, (200, 0, 0))
    second = _upload_via_route(logged_in_admin, pid, (0, 200, 0))
    assert first["ok"] and second["ok"] and first["is_main"]
    with app.app_context():
        db.session.add(ProductVariant(product_id=pid, label="أحمر", image_filename=first["filename"]))
        db.session.commit()
        folder = get_upload_dir("products")
        assert os.path.exists(os.path.join(folder, first["filename"]))

    r = logged_in_admin.post(f"/admin/products/{pid}/images/{first['id']}/replace",
                             data={"image": (io.BytesIO(_jpeg_bytes((0, 0, 200))), "edited.jpg")},
                             content_type="multipart/form-data")
    data = r.get_json()
    assert r.status_code == 200 and data["ok"] and data["relinked_variants"] == 1
    new = data["filename"]
    with app.app_context():
        p = db.session.get(Product, pid)
        row = db.session.get(ProductImage, first["id"])
        assert row.image_filename == new and row.sort_order == 0       # نفس البلاصة
        assert p.image_filename == new                                  # الرئيسية تبعت
        assert ProductVariant.query.filter_by(label="أحمر").one().image_filename == new
        assert db.session.get(ProductImage, second["id"]).image_filename == second["filename"]
        folder = get_upload_dir("products")
        assert os.path.exists(os.path.join(folder, new))
        assert not os.path.exists(os.path.join(folder, first["filename"]))   # القديم تمسح


def test_replace_keeps_old_file_when_shared_with_a_copy(logged_in_admin, app, seed):
    """المنتج المنسوخ يشارك نفس الملفات — ما نمسحوش القديم إذا مازال مستعمل."""
    import io, os
    from models import Product, ProductImage
    from services.uploads import get_upload_dir
    pid = _mk_product(app, "rp-b", 100)
    img = _upload_via_route(logged_in_admin, pid, (90, 90, 90))
    with app.app_context():
        copy = Product(name="نسخة", slug="rp-b-copy", price=100, image_filename=img["filename"])
        db.session.add(copy); db.session.flush()
        db.session.add(ProductImage(product_id=copy.id, image_filename=img["filename"], sort_order=0))
        db.session.commit()
    r = logged_in_admin.post(f"/admin/products/{pid}/images/{img['id']}/replace",
                             data={"image": (io.BytesIO(_jpeg_bytes((1, 2, 3))), "e.jpg")},
                             content_type="multipart/form-data")
    assert r.get_json()["ok"]
    with app.app_context():
        assert os.path.exists(os.path.join(get_upload_dir("products"), img["filename"]))


def test_replace_validation_and_ownership(logged_in_admin, app, seed):
    import io
    a = _mk_product(app, "rp-c", 100); b = _mk_product(app, "rp-d", 100)
    img = _upload_via_route(logged_in_admin, a, (5, 5, 5))
    url = f"/admin/products/{a}/images/{img['id']}/replace"
    assert logged_in_admin.post(url, data={}, content_type="multipart/form-data").status_code == 400
    bad = logged_in_admin.post(url, data={"image": (io.BytesIO(b"not an image"), "x.jpg")},
                               content_type="multipart/form-data")
    assert bad.status_code == 400 and not bad.get_json()["ok"]
    # صورة منتج آخر ← 404
    assert logged_in_admin.post(f"/admin/products/{b}/images/{img['id']}/replace",
                                data={"image": (io.BytesIO(_jpeg_bytes()), "x.jpg")},
                                content_type="multipart/form-data").status_code == 404


def test_replace_requires_login(client, app, seed):
    r = client.post("/admin/products/1/images/1/replace", data={})
    assert r.status_code in (302, 401, 400)


def test_product_form_has_edit_buttons_and_editor_hooks(logged_in_admin, app, seed):
    pid = _mk_product(app, "rp-e", 100)
    _upload_via_route(logged_in_admin, pid, (9, 9, 9))
    html = logged_in_admin.get(f"/admin/products/{pid}/edit").get_data(as_text=True)
    assert "data-edit" in html and "قص / تدوير / قلب" in html
    assert "editFromUrl" in html and "/replace" in html


def test_offer_form_exposes_current_banner_for_editing(logged_in_admin, app, seed):
    from models import Offer
    with app.app_context():
        o = Offer(title="عرض", slug="offer-edit-test", bundle_price=1000, banner_filename="b.webp", is_active=True)
        db.session.add(o); db.session.commit(); oid = o.id
    html = logged_in_admin.get(f"/admin/offers/{oid}/edit").get_data(as_text=True)
    assert 'data-current-image="' in html and "img/offers/b.webp" in html
    # عرض جديد بلا بانر ← بلا زر
    new = logged_in_admin.get("/admin/offers/new").get_data(as_text=True)
    assert "data-current-image" not in new


def test_no_moroccan_or_non_algerian_forms_anywhere_in_project():
    """
    حارس اللهجة: ما نخلّيوش كلمات مغربية/مصرية ترجع للنصوص والتعليقات.
    (الماضي المؤنث الجزائري «-ت»: وصلت، خرجت، رجعت — ماشي «-ات»؛ و«تهزّها» ماشي «تحوّزها».)
    """
    import os, re
    AR = r"\u0621-\u064A"
    forbidden = (
        # مغربي
        "ديال", "ديالي", "ديالك", "ديالو", "دابا", "بغيت", "نبغي", "تبغي", "غادي", "شنو",
        "مزيان", "فين", "دغيا", "واخا", "عفاك", "خاصني", "خاصك", "خاصو",
        # ماضي مؤنث بصيغة «-ات» (جزائري: -ت)
        "وصلات", "وصلاتك", "وصلاتنا", "وصلاتش", "خرجات", "رجعات", "اتبعثات", "تبعثات",
        "تبعثاتش", "اتلغات", "تربطات", "خلصات",
        # «حوّز» مغربي ← «هزّ»
        "تحوّز", "تحوّزها", "يحوّز", "يحوّزها", "نحوّز",
        # مصري/خليجي/شامي
        "دلوقتي", "عشان", "ازاي", "ليش", "شحال", "كده", "لسه", "هيك", "كتير", "منيح",
    )
    pat = re.compile(rf"(?<![{AR}])({'|'.join(forbidden)})(?![{AR}])")
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    hits = []
    for dp, dns, fns in os.walk(root):
        dns[:] = [d for d in dns if d not in ("tests", ".git", "__pycache__", "vendor", "node_modules")]
        for fn in fns:
            if not fn.endswith((".py", ".html", ".js")):
                continue
            path = os.path.join(dp, fn)
            for i, line in enumerate(open(path, encoding="utf-8"), 1):
                m = pat.search(line)
                if m:
                    hits.append(f"{os.path.relpath(path, root)}:{i} «{m.group(0)}»")
    assert not hits, "كلمات غير جزائرية رجعت:\n" + "\n".join(hits[:20])


def test_orders_list_shows_carrier_and_tracking_columns(logged_in_admin, app, seed):
    """قائمة الطلبيات تبيّن «شركة التوصيل» و«رقم التتبع» (— إذا ماكانوش)."""
    with app.app_context():
        db.session.add(Order(
            order_number="BD-COL0001", customer_name="زبون", customer_phone="0555111222",
            wilaya_code="16", wilaya_name="الجزائر", subtotal=500, shipping_price=0,
            total=500, status=Order.STATUS_CONFIRMED, delivery_type="home",
            carrier_name="ZR Express", tracking_number="ZRX-998877",
        ))
        db.session.add(Order(
            order_number="BD-COL0002", customer_name="زبون", customer_phone="0555111333",
            wilaya_code="16", wilaya_name="الجزائر", subtotal=500, shipping_price=0,
            total=500, status=Order.STATUS_PENDING, delivery_type="home",
        ))
        db.session.commit()
    html = logged_in_admin.get("/admin/orders").get_data(as_text=True)
    assert "<th>شركة التوصيل</th><th>رقم التتبع</th>" in html
    assert "ZR Express" in html and "ZRX-998877" in html


def _bordereau_order(app, number, tracking=None):
    with app.app_context():
        o = Order(
            order_number=number, customer_name="دانون عبد الرؤوف",
            customer_phone="0697015018", wilaya_code="43", wilaya_name="ميلة",
            commune="شلغوم العيد", address_note="حي 5 جويلية", subtotal=3000,
            shipping_price=600, total=3600, status=Order.STATUS_CONFIRMED,
            delivery_type="home", carrier_name="ZR Express", tracking_number=tracking,
        )
        db.session.add(o)
        db.session.commit()
        return o.id


def test_bordereau_single_order_shows_label_and_barcode(logged_in_admin, app, seed):
    oid = _bordereau_order(app, "BD-BRD0001", tracking="ZRX-998877")
    resp = logged_in_admin.get(f"/admin/orders/{oid}/bordereau")
    html = resp.get_data(as_text=True)
    assert resp.status_code == 200
    assert "ZRX-998877" in html and "<svg" in html
    assert "دانون عبد الرؤوف" in html and "0697015018" in html
    assert "3,600 دج" in html and "شلغوم العيد" in html


def test_bordereau_without_tracking_falls_back_to_order_number(logged_in_admin, app, seed):
    oid = _bordereau_order(app, "BD-BRD0002")
    html = logged_in_admin.get(f"/admin/orders/{oid}/bordereau").get_data(as_text=True)
    assert "BD-BRD0002" in html and "ما تبعثتش لشركة التوصيل بعد" in html


def test_bordereau_bulk_prints_one_label_per_selected_order(logged_in_admin, app, seed):
    a = _bordereau_order(app, "BD-BRD0003", tracking="T-1")
    b = _bordereau_order(app, "BD-BRD0004", tracking="T-2")
    resp = logged_in_admin.post("/admin/orders/bordereau", data={"order_ids": [a, b]})
    html = resp.get_data(as_text=True)
    assert html.count('<section class="label">') == 2


def test_bordereau_bulk_requires_login_and_selection(app, logged_in_admin, seed):
    empty = logged_in_admin.post("/admin/orders/bordereau", data={})
    assert empty.status_code == 302     # ما اخترت حتى طلبية → رجوع للقائمة
    anon = app.test_client()            # ⚠️ عميل جديد: `client` هو نفسو logged_in_admin
    assert anon.post("/admin/orders/bordereau", data={"order_ids": [1]}).status_code in (302, 401, 403)
    assert anon.get("/admin/orders/1/bordereau").status_code in (302, 401, 403)


def test_orders_list_has_bulk_bordereau_button(logged_in_admin, app, seed):
    _bordereau_order(app, "BD-BRD0005")
    html = logged_in_admin.get("/admin/orders").get_data(as_text=True)
    assert "بوردرو المختارة" in html
