"""الـ Adapter = الملف الوحيد اللي يربط الأداة بمتجر Budac.

⚠ أسماء الجداول والأعمدة تحت هي تخمين مبني على ملف Excel تاع المنتجات.
   لو تختلف عند المتجر، عدّل الـ SQL تاع BudacAdapter برك (3 استعلامات)، والباقي ما يتبدلش.

القاعدة: صور المتجر ما نكوبيوهاش — نرجّعو رابطها العام (و مسارها لو كان في نفس السيرفر).
"""
import os
from pathlib import Path

from . import db
from .config import cfg


class BaseAdapter:
    def search_products(self, q, limit=12):
        """يرجّع [{id, name, price, thumb}]"""
        raise NotImplementedError

    def get_product(self, product_id):
        """يرجّع {id, name, price, url, images:[{id, url, local_path}]} أو None"""
        raise NotImplementedError


def _product_url(slug):
    return cfg.shop_product_url.format(slug=slug)


class BudacAdapter(BaseAdapter):
    # ===== عدّل هنا لو لزم =====
    SQL_SEARCH = """
        SELECT id, slug, name, price FROM products
        WHERE is_active = TRUE AND (%(q)s = '' OR name ILIKE %(like)s OR slug ILIKE %(like)s)
        ORDER BY id DESC LIMIT %(limit)s"""
    SQL_PRODUCT = "SELECT id, slug, name, price FROM products WHERE id = %(id)s"
    SQL_IMAGES = "SELECT filename FROM product_images WHERE product_id = %(id)s ORDER BY position, id"
    # مجلد صور المتجر في السيرفر + بداية رابطها العام
    UPLOADS_DIR = os.getenv("BUDAC_UPLOADS_DIR", "static/uploads")
    UPLOADS_URL = os.getenv("BUDAC_UPLOADS_URL", "/static/uploads")
    # ============================

    def _img(self, filename):
        path = Path(self.UPLOADS_DIR) / filename
        url = filename if filename.startswith("http") else f"{cfg.public_base_url}{self.UPLOADS_URL}/{filename}"
        return {"id": filename, "url": url, "local_path": str(path.resolve()) if path.exists() else None}

    def search_products(self, q, limit=12):
        with db.cursor(commit=False) as cur:
            cur.execute(self.SQL_SEARCH, {"q": q, "like": f"%{q}%", "limit": limit})
            rows = cur.fetchall()
            out = []
            for r in rows:
                cur.execute(self.SQL_IMAGES + " LIMIT 1", {"id": r["id"]})
                first = cur.fetchone()
                out.append({"id": str(r["id"]), "name": r["name"], "price": r["price"],
                            "thumb": self._img(first["filename"])["url"] if first else None})
        return out

    def get_product(self, product_id):
        with db.cursor(commit=False) as cur:
            cur.execute(self.SQL_PRODUCT, {"id": product_id})
            p = cur.fetchone()
            if not p:
                return None
            cur.execute(self.SQL_IMAGES, {"id": product_id})
            imgs = [self._img(r["filename"]) for r in cur.fetchall()]
        return {"id": str(p["id"]), "name": p["name"], "price": p["price"],
                "url": _product_url(p["slug"]), "images": imgs}


class ExcelDemoAdapter(BaseAdapter):
    """للتجربة المحلية برك: يقرا ملف Excel تاع المنتجات (بلا صور، الصور تترفع يدويًا)."""

    def __init__(self, xlsx_path):
        from openpyxl import load_workbook
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        ws = wb["المنتجات"]
        rows = list(ws.iter_rows(values_only=True))
        head = {name: i for i, name in enumerate(rows[0])}
        self.items = []
        for r in rows[1:]:
            if r[head["المعرّف"]] is None:
                continue
            self.items.append({"id": str(r[head["المعرّف"]]), "slug": r[head["الرابط (slug)"]],
                               "name": r[head["اسم المنتج"]], "price": r[head["السعر"]]})

    def search_products(self, q, limit=12):
        hits = [i for i in self.items if q.lower() in (i["name"] or "").lower()][:limit]
        return [{"id": i["id"], "name": i["name"], "price": i["price"], "thumb": None} for i in hits]

    def get_product(self, product_id):
        for i in self.items:
            if i["id"] == str(product_id):
                return {"id": i["id"], "name": i["name"], "price": i["price"],
                        "url": _product_url(i["slug"]), "images": []}
        return None
